from pathlib import Path
import json, shutil, re, zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT=Path(__file__).resolve().parent
records=json.loads((OUT/'dataset_registry.json').read_text())['records']
byid={r['id']:r for r in records}

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.6)
sec.left_margin=sec.right_margin=Inches(.7)
styles=doc.styles
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    styles[n].font.name='Arial';styles[n].font.color.rgb=RGBColor(0,0,0)
    rf=styles[n]._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in list(rf.attrib):
        if 'Theme' in attr or 'theme' in attr: del rf.attrib[attr]
    rf.set(qn('w:ascii'),'Arial');rf.set(qn('w:hAnsi'),'Arial')
styles['Normal'].font.size=Pt(10)
styles['Normal'].paragraph_format.space_after=Pt(6)
styles['Normal'].paragraph_format.line_spacing=1.08
styles['Title'].font.size=Pt(23)
styles['Heading 1'].font.size=Pt(16)
styles['Heading 1'].paragraph_format.space_before=Pt(12)
styles['Heading 1'].paragraph_format.space_after=Pt(7)
styles['Heading 2'].font.size=Pt(12)
styles['Heading 2'].paragraph_format.space_before=Pt(9)
styles['Heading 2'].paragraph_format.space_after=Pt(4)
for st in styles:
    for el in list(st._element.iter(qn('w:pBdr'))):el.getparent().remove(el)
header=sec.header.paragraphs[0];header.text='Kevin paper dementia MRI';header.style='Normal';header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Research preparation  •  11 October 2026  •  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for r in footer.runs:r.font.size=Pt(8)
doc.core_properties.title='Kevin paper datasets and training plan'
doc.core_properties.subject='Dementia MRI dataset acquisition and proposed training experiment'
doc.core_properties.author='Arman Kazi'

def p(text='',style=None):return doc.add_paragraph(text,style)
def h(text,level=1):return doc.add_heading(text,level)
def page(title):doc.add_page_break();h(title)
def link(par,label,url):
    rel=par.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    a=OxmlElement('w:hyperlink');a.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'174A74');pr.append(c);r.append(pr)
    t=OxmlElement('w:t');t.text=label;r.append(t);a.append(r);par._p.append(a)
def source(label,url):
    q=p();link(q,label,url);q.paragraph_format.space_after=Pt(3)
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for vals in rows:
        cells=t.add_row().cells
        for c,txt in zip(cells,vals):
            if isinstance(txt,tuple):link(c.paragraphs[0],txt[0],txt[1])
            else:c.text=str(txt)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        for j,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[j])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE8EF' if i==0 else 'FFFFFF');tc.append(sh)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tc.append(borders);mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');mar.append(e)
            tc.append(mar)
            for q in c.paragraphs:
                q.paragraph_format.space_after=Pt(2);q.paragraph_format.space_before=Pt(2);q.paragraph_format.line_spacing=1.03
                for r in q.runs:r.font.size=Pt(9);r.bold=(i==0)
    p();return t

p('Kevin paper datasets and training plan','Title')
p('Prepared for Arman and Kevin  11 October 2026','Subtitle')
p('The immediate priority is a reproducible dementia experiment with documented MRI sources, patient grouping, genuine loss curves and a fair comparison with the current model. The dataset search covers 35 entries, including overlapping releases and data streams. Two open cohorts now provide a practical starting point; larger cohorts require access applications.')
h('What the recording requests',2)
table(['Recording time','Kevin paper requirement'],[
('01:03–01:31','Find as many suitable dementia MRI databases as possible; prefer age and clinical information.'),
('01:31–02:13','Build a new training dataset and regenerate results while preserving the existing 100-case radiologist component.'),
('02:13–03:18','Rewrite methods for the actual databases used; regenerate figures and statistics; prioritize published NIH or NCBI-linked sources.')],[1.25,5.85])
p('The STS section ends at approximately 01:03 and is outside this work. Later chat instructions add model improvement, a 100-case comparison with the old model and neuroradiologists, and long training runs on a virtual computer the user will provide. A 100-epoch maximum is planned; validation loss controls checkpoint selection and early stopping.')
h('Data acquired and checked',2)
table(['Dataset','People and original T1 volumes','Immediate role'],[
('Besga et al. via NCBI / Zenodo','78: 35 AD, 19 controls, 24 bipolar','Proposed AD/control pool of 54; bipolar retained separately.'),
('OpenNeuro ds007561 v1.0.0','20: 1 AD, 5 MCI, 14 controls','Reserve entire source for exploratory external evaluation; MCI retained for descriptive predictions only.')],[2.0,2.7,2.4])
p('The acquisition contains 98 source participant records and 98 original T1 volumes across two cohorts; cross-repository identity independence has not been verified. Of these, 74 participants were preprocessed into 666 slices: 54 Besga AD/control participants and all 20 OpenNeuro participants. The 24 bipolar participants remain outside the primary endpoint. The clinical split contains only 32 training, 11 validation and 11 internal-test participants. OpenNeuro remains an external pilot of 15 AD/control participants, including one AD case, plus five MCI participants for descriptive predictions. Slices, processed versions and augmentations never increase participant counts; no four-stage labels were invented.')
p('For Besga, the approved acquired demographics table supplies source diagnosis code, study ID, age and gender code. Thirteen selected CSVs comprise that table and biomarker files. Although the paper describes broader cognitive and clinical scales, the provider says clinical_data_corrected.csv can be ignored; it was not extracted or used. NPI/FAST and other detailed scales are therefore not claimed as verified usable acquired fields, and biomarker linkage remains unaudited.')
source('Besga dataset paper in NCBI PMC','https://pmc.ncbi.nlm.nih.gov/articles/PMC7658647/')
source('OpenNeuro release and public clinical metadata','https://openneuro.org/datasets/ds007561/versions/1.0.0')

page('Source quality and access')
p('A reputable source means an identifiable study team, an original dataset publication or official institutional provider, explicit imaging and reference-label descriptions, usable participant identifiers, and documented use terms. PubMed indexing is a route to the paper, not NIH endorsement of a model or permission to download its scans.')
table(['Resource','Evidence and access','Decision'],[
('Besga / Zenodo','Original data paper in PMC; repository CC BY 4.0; linked original T1 and clinical records.','Acquired original scans. Subject-level splits and anatomical review precede training.'),
('OpenNeuro ds007561','Public CC0 release; pinned commit and published git-annex hashes.','Acquired 20 original T1 scans. One AD case makes it a pilot only.'),
('ADNI','NIH-supported longitudinal imaging and clinical study; LONI investigator application.','High priority once access is available.'),
('OASIS','Official Washington University datasets; reviewed requests and dataset-specific access terms.','Choose a suitable release. OASIS 1, 2 and 3 overlap.'),
('AIBL and NACC','Established provider programs with clinical and imaging data; applications and agreements.','Strong expansion candidates; obtain MRI-specific participant counts.'),
('ANMerge and BrainLat','Original published patient cohorts; Synapse access controls.','Useful MRI and clinical expansion after permission.')],[1.25,3.55,2.3])
p('The existing 44,000-image Kaggle dementia collection remains a legacy image benchmark. It was augmented before this project received it and lacks patient and original-family IDs. More Kaggle mirrors cannot establish new independent patient evidence. The locally acquired Mendeley cohort has 24 subject folders and T2/FLAIR scans; it is not automatically compatible with T1-based training.')
h('Datasets from Abid’s email',2)
p('Abid’s 11 October email lists ADNI, OASIS-1, OASIS-2, OASIS-3, OASIS-4, NACC, AIBL, MIRIAD, NIFD/FTLDNI and ALLFTD. All ten are included in this registry and now explicitly tracked in abid_dataset_crosswalk.json. They are acquisition candidates, not scans already acquired or authorized. Prioritize MRI-linked clinical severity records; audit actual class counts, dates and access before adding a cohort. The provider advises against pooling OASIS-1, -2 and -3 because participants can overlap under different IDs. The four-stage experiment will need a separate frozen revision if suitable CDR-labeled cohorts become available.')
h('What needs to happen for larger cohorts',2)
p('Kevin or the supervising investigator will need to supply actual affiliation, investigator identity, proposed research use and any provider-required ethics or institutional documentation. The package does not invent these details or submit applications. Data access may also restrict cloud environments, redistribution or AI-tool use; the approved processing location must match the actual agreement.')
p('Proposed access statement: We propose research on dementia MRI classification using subject-linked structural MRI and accompanying clinical reference labels. We will preserve participant grouping, audit source overlap, compare validation-selected models on untouched cohorts, and report sample counts, uncertainty and failures. No clinical decisions or re-identification are planned. Data will remain within the provider-approved environment, with required citations and acknowledgements.')
source('ADNI official data access','https://adni.loni.usc.edu/data-samples/adni-data/')
source('OASIS official access and cohort information','https://sites.wustl.edu/oasisbrains/')

# Compact inventory with links to the most relevant original/provider source.
notes={
'ADNI':'Development priority after approval; phases share participants.',
'OASIS-1':'416 participants; legacy alternative to OASIS 2 or 3.',
'OASIS-2':'150 participants and 373 sessions; overlaps OASIS 1 and 3.',
'OASIS-3':'Longitudinal T1 and clinical data; displayed 1,378/1,379 total is inconsistent.',
'OASIS-4':'Clinical memory cohort; official FAQ says separate from OASIS 1–3.',
'AIBL':'AD/MCI/control imaging and clinical data; request approved subset.',
'NACC_MIXED_PROTOCOL':'Clinical MRI collection; shared-center participants may overlap other cohorts.',
'NACC_SCAN':'Standardized MRI stream within NACC; not an independent extra cohort.',
'MIRIAD':'69 people and 708 T1 scans; AD/control with age and MMSE.',
'BRAINLAT':'AD and other diagnoses with clinical data; Synapse downloads require access.',
'SYNAPSE_OPENNEURO_DS007561':'20 T1 participants acquired; 1 AD, 5 MCI, 14 controls.',
'BESGA_ZENODO_3935636':'78 original T1 participants acquired; 54 AD/control candidates.',
'DELCODE':'~1,000 enrolled; controls, SCD, MCI and mild dementia; MRI subset to verify.',
'EPAD':'Over 2,000 enrolled without dementia at baseline; conditional preclinical use.',
'ALLFTD':'FTD spectrum; >3,000 clinical records does not mean that many MRI scans.',
'GENFI':'Genetic FTD families; auxiliary disease task, not AD staging.',
'FTLDNI_NIFD':'FTD imaging and clinical data; LONI access request.',
'4RTNI':'Four-repeat tauopathy spectrum; auxiliary differential-diagnosis task.',
'ADC':'Amsterdam memory-clinic cohort; collaborator-mediated request.',
'ABCDS':'Down syndrome and AD; distinct population and disease trajectory.',
'MEMENTO':'2,323 enrolled; cognitive complaints/MCI; application and secure analysis.',
'BIOFIND':'324 reported; check matched MRI fields and DPUK release.',
'NIMROD':'167 in DPUK matrix; inflammatory/dementia research; verify released subset.',
'PREVENT_DEMENTIA':'Midlife risk cohort; distinct from PREVENT AD; conditional use.',
'AMPLE':'Lewy body dementia study; MRI subset and release need confirmation.',
'QMIN_MC':'NHS memory-clinic MRI; published and portal counts differ.',
'UK_BIOBANK':'Large population imaging study; new applications currently paused.',
'ANMERGE':'1,702 clinical participants, 453 with MRI measures; raw T1 available by request.',
'ARWIBO':'AD/MCI/control structural MRI; NeuGRID2 agreement.',
'EDSD':'471 MPRAGE and 471 DTI scans; these are not 942 people.',
'I_ADNI':'262 participants; keep AD and FTD labels distinct.',
'PHARMACOG':'147 MCI participants over seven timepoints; not baseline dementia.',
'NIGERIAN_MRI':'88 participants, mixed clinical sequences/diagnoses; signed DUA.',
'MENDELEY_KCJT4V658X':'Already local: 24 folders and 474 DICOMs; T2/FLAIR; review identity fields.',
'PREVENT_AD':'308 listed in CONP V2; cognitively healthy at-risk adults; account needed.'}
ordered=['BESGA_ZENODO_3935636','SYNAPSE_OPENNEURO_DS007561','ADNI','OASIS-3','OASIS-4','AIBL','NACC_MIXED_PROTOCOL','NACC_SCAN','MIRIAD','ANMERGE','BRAINLAT','DELCODE','ARWIBO','EDSD','I_ADNI','OASIS-1','OASIS-2','NIGERIAN_MRI','MENDELEY_KCJT4V658X','ADC','MEMENTO','BIOFIND','NIMROD','QMIN_MC','EPAD','PHARMACOG','PREVENT_AD','PREVENT_DEMENTIA','ALLFTD','GENFI','FTLDNI_NIFD','4RTNI','ABCDS','AMPLE','UK_BIOBANK']
for part, ids in enumerate([ordered[:12],ordered[12:24],ordered[24:]],1):
    page('Dataset inventory '+str(part))
    p('Names link to original publications or official providers. The machine-readable registry contains more detail, all source links, access conditions and overlap notes. Entries are candidates, not a claim that all are acquired, suitable, or independent.')
    rows=[]
    for rid in ids:
        r=byid[rid];name=r['name']
        if rid=='SYNAPSE_OPENNEURO_DS007561':name='OpenNeuro ds007561'
        if rid=='BESGA_ZENODO_3935636':name='Besga AD and controls'
        if len(name)>57:name=rid.replace('_',' ')
        access='Local' if rid in ['BESGA_ZENODO_3935636','SYNAPSE_OPENNEURO_DS007561','MENDELEY_KCJT4V658X'] else ('Paused' if rid=='UK_BIOBANK' else 'Request / terms')
        rows.append(((name,r['sources'][0]),access,notes[rid]))
    table(['Dataset or release','Access now','Suitability and counting boundary'],rows,[1.75,1.05,4.3])
    if part==3:
        p('FTD, Lewy body, Down syndrome, preclinical risk and healthy-aging cohorts must not be relabeled as AD severity stages. Imaging portals such as DPUK, NITRC, LONI, NCBI and AD Workbench are access or discovery platforms, not additional independent cohorts. Ex-vivo specimens, EEG-only sets, synthetic MRI and derived-only tables were not selected for the primary clinical MRI experiment.')

page('Proposed methods for the new experiment')
p('This is a prospective methods draft. Training has not yet been performed for the new paper, and model performance must be filled from actual runs. The long runs will begin on Runpod after account authentication, resource selection and a CUDA smoke test.')
h('Study endpoints and models',2)
p('Retain a four-class head for the provided legacy labels MildDemented, ModerateDemented, NonDemented and VeryMildDemented. Add a separate AD-versus-control head for the newly acquired clinical labels. Clinical AD, MCI and dementia severity are different targets; AD and MCI will not be assigned invented stage labels. The bipolar participants will be excluded from the primary AD/control training objective, and MCI will be retained for separately described predictions.')
p('Compare a stage-only model with a model that also learns from the clinical AD/control data. Both begin from the same ImageNet initialization rather than the current study checkpoint. This avoids carrying the old 100-case training exposure into a supposedly fresh model. The shared architecture and fixed training choices isolate the effect of added clinical data; any benefit remains a hypothesis until evaluated.')
h('Splitting and preprocessing',2)
p('Use the frozen perceptual-grouped dementia manifest. The preparation removes 1,863 images in exact-pixel or identical-dHash groups linked to the 100-case set and 14,959 images from mixed-label dHash groups, leaving 27,178 images (19,023 training, 2,705 validation, 5,450 test). These coarse-hash exclusions are conservative: a shared hash does not prove a shared patient or erroneous label. All four classes remain represented. Group all clinical scans and derived slices by participant before splitting. Besga AD/control participants form a seeded development and internal-test pool. Keep the whole OpenNeuro source outside model fitting and checkpoint selection. Shared participants across repositories remain a provenance question beyond pixel hashing.')
p('The frozen clinical preprocessing uses affine-based RAS orientation, 1.5 mm resampling, per-volume intensity normalization and nine axial slices at fixed 45–85% fractions of the detected cranial extent. It records the source volume, orientation, slice locations and processing values. This is a heuristic adaptation to a 2D model, without validated skull stripping, atlas registration or scanner harmonization. Gross visual review found plausible rendered anatomy; it is not radiologist or clinical quality assurance.')
p('The initial 20–80% slice range was revised after gross anatomical inspection showed excessive inferior tissue. Initial QC sheets included diagnosis labels, and both sources were inspected before this final rule was frozen; no model outcomes were viewed and no subject was removed on that basis. OpenNeuro is excluded from model fitting and selection, but it was not unseen during preprocessing preparation. Report it as an exploratory external pilot, and preserve this adaptation history.')
p('Age and source gender remain descriptive metadata, not image-model inputs. Besga AD participants have mean age 79.1 years, versus 73.6 for controls. Images can still encode age, so report age distributions and assess age/source confounding before interpreting any improvement as disease-specific. Cross-source appearance and unknown participant overlap also remain limitations.')
h('Training and selection',2)
p('Configure up to 100 full training epochs per condition, validation-based early stopping, resumable checkpoints, random seeds and complete optimizer settings. Save unaltered epoch loss and accuracy values for every head and source. Use validation performance to select checkpoints and any calibration parameters. Test data and the 100-case set must not determine architecture, epochs, preprocessing, hyperparameters or checkpoint selection.')
p('Use participant-level aggregation for the clinical endpoint so nine slices do not count as nine people. For the legacy comparison, evaluate the 1,058-image intersection held out in both the old strict manifest and the new experiment. Unknown legacy patient identities still limit independence. Report the results for each endpoint and source separately before discussing an overall improvement.')

page('Results and figures to regenerate')
table(['Output','Required content'],[
('Cohort flow and demographics','Source versions; requested, downloaded and eligible patients; scans; exclusions; split counts; class balance; age/sex/site and missingness.'),
('Training curves','Raw train and validation loss and accuracy for every epoch, head and condition; selected epoch, learning rate and early-stop point. No cosmetic editing of results.'),
('Discrimination','Per-class sensitivity, specificity, precision, F1; macro F1 and balanced accuracy; confusion matrices; ROC and precision-recall curves where supported.'),
('Calibration','NLL, Brier score, ECE, reliability plots and confidence distributions; calibration fit only to validation predictions.'),
('Generalization','Separate legacy internal, clinical internal and external pilot results. Patient-level uncertainty; report absent classes and very small class counts.'),
('Model and radiologist comparison','One case-aligned CSV for all 100 images: reference, old model, new model, probabilities, reader labels, disagreement and missing-reader status.'),
('Explanation analysis','Representative Grad-CAM successes and failures, reproducibility and prediction invariance. Heatmaps are exploratory attribution, not proven anatomy.')],[1.7,5.4])
h('How to decide whether the model is stronger',2)
p('Require a prespecified improvement on untouched evaluation data without unacceptable class-specific regressions. Assess balanced accuracy and macro F1 alongside calibration, not training accuracy alone. Repeat final experiments across independent seeds if GPU resources allow; include every run in reporting. Low training loss with rising validation loss is overfitting, not evidence of a stronger model.')
p('The current checkpoint is already 100/100 on the historical 100 images. Those cases match 66 training, 9 validation and 25 test images in its strict manifest. A new model cannot exceed 100% accuracy there. Preserve and run the requested comparison after model selection, but label it a reused image benchmark; it cannot establish independent superiority over neuroradiologists.')
p('The OpenNeuro pilot contains only one AD case. Its uncertainty will be large, and it cannot substantiate broad clinical superiority. A larger approved cohort remains necessary for convincing external evaluation.')

page('Execution status and handoff')
h('Completed',2)
p('Recording requirements were extracted; 35 candidate entries were checked against provider or original publication sources; the current local data were audited; two open MRI sources were acquired; and a manifest staging checker was built. Acquisition summaries record the actual integrity checks and remaining anatomical/provenance limitations. Raw data remain in Git-ignored local folders.')
h('Prepared for the virtual computer',2)
p('The stage-only and multitask MPS smoke tests passed, including checkpoint resume and isolation guards. The reader loader verified three readers with 100 recognized labels each in the bundled runtime. The current model remains unchanged. CUDA preflight and full training are still pending; smoke curves are not paper results. Launch instructions specify data, seeds, epoch budgets and evaluation commands.')
p('Before connecting, obtain the host address, login method, GPU model and VRAM, operating system, persistent disk location, available storage and shutdown policy. Verify CUDA access, disk capacity and package versions on the actual host. Transfer only the chosen local research dataset, without credentials or unrelated files. Provider-controlled datasets require a processing environment allowed by their agreements.')
h('Runpod connection',2)
p('The official Runpod marketplace and version 1.7.1 plugin are installed and enabled; its hosted MCP server is configured. Browser OAuth succeeded and the CLI confirms authenticated status. The current chat has not loaded the new Runpod tools yet; restart Codex and verify the account by listing Pods. No paid compute has been provisioned; launch awaits the user’s payment update. The onboarding guide was followed without installing CLI/Flash or creating API keys.')
h('Remaining work',2)
p('Reload Runpod tools, verify the account and GPU environment, run the long training conditions, and evaluate the selected checkpoints once. Then regenerate the result tables, curves, 100-case comparison and manuscript text. Better performance is an experimental objective, not a result established by this packet.')
h('Files to use',2)
table(['File','Purpose'],[
('dataset_registry.json','All 35 candidate entries, sources, metadata and access constraints.'),
('besga_acquisition.json and openneuro_acquisition.json','Aggregate evidence for the two new acquisitions.'),
('current_local_audit.json','Legacy dataset and model state, with known exposure limitations.'),
('cohort_manifest_template.csv and data_dictionary.json','Source-scan schema with explicit labels, patient groups and pending states.'),
('validate_cohort_manifest.py','Read-only structural checks. An empty or pending manifest fails; passing is not clinical validation.'),
('recording_kevin_excerpt.txt','Machine transcription of the relevant recording section.'),
('training/ and runpod_setup_status.json','Run scripts, evaluation, transfer preparation and verified connection status.'),
('abid_dataset_crosswalk.json','Tracks all ten datasets from the emailed assessment and their access/fit.')],[3.25,3.85])
doc.save(OUT/'Kevin_Paper_Datasets_and_Training_Plan.docx')
print(json.dumps({'docx':str(OUT/'Kevin_Paper_Datasets_and_Training_Plan.docx'),'registry_entries':len(records),'source_links':sum(len(r['sources']) for r in records)},indent=2))
