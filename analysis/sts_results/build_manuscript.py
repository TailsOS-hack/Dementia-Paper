"""Build the Results manuscript from audited frozen study artifacts.

Run the three analysis modules first. Uses python-docx; no model inference.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

STUDY = Path('docs/kevin_paper_20261011/publication')


def make_content(root: Path, source_commit: str):
    study = root / STUDY
    selection = json.loads((study/'results/validation_selection.json').read_text())
    replay = json.loads((study/'figures/clean_replay/replay_metadata.json').read_text())
    histories = {}
    for key, c in selection['candidates'].items():
        with (study/f"training/{c['mode']}_seed{c['seed']}/history.csv").open() as f:
            histories[key] = list(csv.DictReader(f))
    def label(key):
        c = selection['candidates'][key]
        return f"{'Stage' if c['mode']=='stage_only' else 'Multitask'} {str(c['seed'])[-2:]}"
    blocks=[]
    def add(kind, **kw): blocks.append(dict(kind=kind, **kw))
    def p(text): add('paragraph', text=text)
    def h(text): add('heading', text=text)
    def ref(label_, relative):
        add('reference', text=label_, url=f'https://github.com/TailsOS-hack/Dementia-Paper/blob/{source_commit}/{STUDY.as_posix()}/{relative}')
    add('title', text='Dementia MRI Training and Benchmark Results')
    add('byline', text='Arman Kazi | STS manuscript Results | October 18 2026')
    p('Six MobileNetV3-Large training runs achieved 99.32% to 99.89% accuracy on the complete 5,450-image test split. The runs comprise two conditions with three seeds each and are archived as the October 11 Round 1 study. The full-test replay is descriptive and post hoc; the 1,058-image matched comparison retains its prespecified primary role. These image-level results do not establish patient-independent or clinical performance.')
    h('Training and validation selection')
    p('Stage-only learned the four supplied image labels; multitask added a separate AD/control auxiliary head with a shared backbone. Both conditions used seeds 20261011, 20261012 and 20261013, completing 305 epochs in total. All six runs stopped under the 15-epoch patience rule before the 100-epoch cap. Checkpoints were selected by validation negative log-likelihood with a required improvement greater than 0.0001. The validation-only preference was stage-only, represented by seed 20261012 at epoch 37; no test result was used to change this preference.')
    rows=[]
    for key,c in selection['candidates'].items():
        rows.append([key.replace('candidate_', 'C'), 'Stage only' if c['mode']=='stage_only' else 'Multitask', str(c['seed']), str(len(histories[key])),str(c['epoch']),f"{c['stage_validation_nll']:.6f}"])
    add('table', caption='Table 1. Completed training and validation-selected checkpoints.', headers=['Run','Condition','Seed','Completed','Selected','Val NLL'],rows=rows,widths=[.42,1.02,1.03,.82,.76,.95])
    add('image', path='results/sts/training/completed_vs_selected_epochs.png', width=5.95, caption='Figure 1. Actual completed epochs and selected checkpoint epochs for the six runs. These are two conditions across three seeds, not six architectures. C1 to C6 correspond to Table 1.')
    add('break')
    h('Complete test split of 5450 images')
    p('The legacy collection contained 44,000 upstream-augmented JPEG images. Group-based exclusions removed 1,863 rows connected to the historical 100-image benchmark and 14,959 rows in mixed-label similarity groups, leaving 27,178 images: 19,023 training, 2,705 validation and 5,450 test. The exclusions do not establish that the supplied labels were wrong. Exact-pixel and identical-dHash connected groups were kept within one split; patient and original augmentation-family identities remain unavailable.')
    p('The complete test split contained 1,342 MildDemented, 1,538 ModerateDemented, 1,324 NonDemented and 1,246 VeryMildDemented images. Table 2 was recomputed from each frozen checkpoint\u2019s saved full-test prediction CSV and checked against the archived replay metrics. No retraining, checkpoint reselection or new inference was used to recompute these historical results.')
    rows=[]
    for key in selection['candidates']:
        m=replay['checkpoints'][key]['splits']['full_new_test']['metrics']
        correct=sum(m['confusion_matrix'][i][i] for i in range(4))
        rows.append([label(key),f'{correct:,}/5,450',str(5450-correct),f"{100*m['accuracy']:.2f}",f"{100*m['balanced_accuracy']:.2f}",f"{100*m['macro_f1']:.2f}"])
    add('table',caption='Table 2. Complete test split performance for all six selected checkpoints.',headers=['Model','Correct / n','Errors','Acc %','BA %','F1 %'],rows=rows,widths=[1.12,1.22,.56,.71,.71,.68])
    p('Acc denotes accuracy; BA is the mean recall across the four supplied classes; F1 is macro-averaged F1. Model suffixes 11, 12 and 13 denote the final two digits of the full seeds in Table 1. Stage-only accuracies were 99.78%, 99.87% and 99.89%; multitask accuracies were 99.32%, 99.72% and 99.60%. Across all six models, balanced accuracy ranged from 99.27% to 99.88% and macro F1 from 99.29% to 99.88%.')
    p('The validation-preferred representative, Stage 12, correctly classified 5,443/5,450 images, with 99.87% accuracy, 99.87% balanced accuracy and 99.87% macro F1. Stage 13 had one fewer observed error, but this post-hoc test difference does not replace the validation-selected representative.')
    h('Relationship to the matched comparison')
    p('The replay protocol explicitly identifies the full 5,450-image evaluation as a post-hoc supplement that contains the primary 1,058-image subset. It therefore supplies a broader descriptive assessment of the six new checkpoints, not an additional independent cohort. Under the old baseline split, these 5,450 images comprise 3,860 training, 532 validation and 1,058 test rows; a baseline comparison on all 5,450 would mix previously seen and held-out images.')
    p('The historical full-test replay did not provide paired uncertainty estimates for this complete cohort. Table 2 therefore reports point estimates only. Detailed per-class counts and all confusion matrices accompany the reproducible analysis in the repository.')
    add('break')
    h('Matched comparison and sensitivity analysis')
    p('The prespecified matched comparison used the same 1,058 exact image rows for the old baseline and all six new checkpoints, representing 966 prepared similarity groups. Of those images, 302 had a similarity-group mate in old training or validation. Excluding those groups yielded the prespecified 756-image sensitivity subset in 733 groups. The smaller subset changes sample composition and does not establish patient independence.')
    # The frozen point estimates are independently checked by supporting.py.
    cohorts=['internal_same_case','old_hash_group_unexposed_sensitivity','original100']
    counts={
      'old':[1057,755,100], 'candidate_1':[1055,753,99], 'candidate_2':[1054,752,100],
      'candidate_3':[1057,755,100], 'candidate_4':[1056,754,99], 'candidate_5':[1057,755,100],
      'candidate_6':[1055,754,99], 'stage_ensemble':[1058,756,100], 'multitask_ensemble':[1056,754,100]}
    # Values in the presentation table are compared to recomputed metrics before export.
    audited=list(csv.DictReader((root/'results/sts/supporting/metrics.csv').open()))
    rows=[]
    for key,values in counts.items():
        name={'old':'Old baseline','stage_ensemble':'Stage ensemble','multitask_ensemble':'Multitask ensemble'}.get(key, label(key) if key.startswith('candidate_') else key)
        rows.append([name]+[f'{v}/{n} ({100*v/n:.2f}%)' for v,n in zip(values,[1058,756,100])])
    add('table', caption='Table 3. Correct classifications and accuracy on the matched and historical benchmarks.', headers=['Model','Matched 1058','Sensitivity 756','Reused 100'],rows=rows,widths=[1.30,1.28,1.28,1.14])
    p('Individual-model accuracy ranged from 99.62% to 99.91% in the matched comparison and 99.47% to 99.87% in the sensitivity analysis. The stage-only ensemble averaged the three frozen probability vectors uniformly and achieved 1,058/1,058 and 756/756 correct classifications. The multitask stage ensemble achieved 1,056/1,058 and 754/756. Ensemble membership, weights and thresholds were not fitted to these test results.')
    p('The frozen analysis used 2,000 paired percentile bootstrap resamples of similarity groups. For the validation-preferred Stage 12 model versus the baseline, the matched accuracy difference was 0.00 percentage points (95% interval -0.29 to +0.28); the sensitivity difference was 0.00 (-0.40 to +0.40). These estimates do not demonstrate a performance improvement. Intervals are exploratory, unadjusted for multiplicity and conditional on the available images and grouping; they exclude uncertainty from training and source selection.')
    h('Reused historical benchmark')
    p('The old baseline, Stage 12, Stage 13, Multitask 11 and both ensembles correctly classified all 100 historical benchmark images; the other three single checkpoints scored 99/100. Of these 100 images, 66 match old training rows, nine old validation rows and 25 old test rows. This is reused descriptive evidence. Perfect observed accuracy and degenerate bootstrap intervals do not imply population certainty or clinical superiority.')
    add('break')
    h('Benchmark comparison and original figures')
    add('image', path='results/sts/supporting/accuracy_comparison.png', width=6.05, caption='Figure 2. Accuracy for the exact-row matched comparison, its similarity-exclusion subset and the reused historical benchmark. The sensitivity subset overlaps the matched cohort, and the historical images are reused. The narrowed horizontal scale makes small differences visible; Table 3 gives the counts and denominators.')
    h('Interpretation')
    p('The results show high agreement with four supplied image labels within this legacy dataset. Upstream augmentation, unknown patient identities and unverified augmentation-family relationships limit generalization claims. The four labels are not a validated mapping to clinical CDR stages. Similarity exclusion cannot remove unknown relationships or prove that every image was equally unseen by the old and new models. Clinical AD/control results require separate endpoints and participant-level evidence; they are not inferred from these four-class percentages.')
    p('The earlier specialist histories are separate from the six October 11 runs summarized here. The 305-epoch total must not be attributed to six architectures, to 18 runs, or to the older specialist checkpoint. The manuscript emphasizes the six-run development study and complete test split while retaining the original primary-versus-supplementary analysis roles.')
    h('Source records and original GitHub figures')
    ref('Training histories and completed Methods', 'reports/Completed_Methods.md')
    ref('Frozen validation selection and analysis protocol', 'results/validation_selection.json')
    ref('Full-test replay protocol and cohort role', 'figures/clean_replay/protocol_before_inference.json')
    ref('Original full-test confusion matrices', 'figures/quantitative/full_new_test_confusion_all_six.png')
    ref('Original six-run loss curves', 'figures/quantitative/stage_loss_all_six_runs.png')
    ref('Original paired accuracy intervals', 'figures/quantitative/paired_accuracy_differences.png')
    ref('Original two-condition architecture figure', 'figures/quantitative/dementia_two_condition_architecture.png')
    ref('Complete matched statistics and historical benchmark', 'reports/Full_Results.md')
    return blocks, counts, audited


def make_retraining_appendix(report_path: Path, *, allow_synthetic: bool = False):
    """Append two pages using a completed retraining report and its own figures.

    Historical content remains unchanged. Synthetic input is rejected by the CLI;
    the internal opt-in exists only for explicitly labelled local fixture tests.
    """
    report_path = Path(report_path).resolve()
    report = json.loads(report_path.read_text())
    synthetic = report.get('synthetic_fixture') is True
    expected_status = 'verified_synthetic_fixture' if synthetic else 'verified_completed_retraining'
    if report.get('study') != 'sts_retrain_20261018' or report.get('status') != expected_status:
        raise ValueError('Appendix requires a validated completed October 18 retraining report')
    if synthetic and not allow_synthetic:
        raise ValueError('Synthetic results cannot be appended to the research manuscript')
    if report.get('completed_runs') != 6 or report.get('heldout_performance_estimated') is not False:
        raise ValueError('Retraining report must contain six completed runs and no held-out performance estimates')
    rows = report['runs']
    names = {f'{mode}_seed{seed}' for mode in ('stage_only', 'multitask') for seed in (20261011, 20261012, 20261013)}
    if len(rows) != 6 or {r['run'] for r in rows} != names:
        raise ValueError('Retraining report has duplicate or missing condition/seed rows')
    if sum(r['completed_epochs'] for r in rows) != report['total_epochs']:
        raise ValueError('Retraining epoch total does not match the six runs')
    summary_path = report_path.parent/'summary.csv'
    with summary_path.open(newline='') as stream:
        summary_rows = list(csv.DictReader(stream))
    summaries = {r['run']: r for r in summary_rows}
    if len(summary_rows) != 6 or set(summaries) != names:
        raise ValueError('Retraining summary CSV has missing or duplicate run rows')
    for row in rows:
        if row['run'] != f'{row["mode"]}_seed{row["seed"]}':
            raise ValueError('Retraining condition/seed does not match its run identity')
        for key in ('completed_epochs', 'selected_epoch', 'selected_train_nll', 'selected_validation_nll',
                    'train_images', 'validation_images'):
            if not math.isfinite(float(row[key])) or not math.isclose(float(row[key]), float(summaries[row['run']][key]),
                                                                    rel_tol=0, abs_tol=1e-11):
                raise ValueError(f'Retraining report/summary mismatch: {row["run"]} {key}')
        if not 1 <= row['selected_epoch'] <= row['completed_epochs']:
            raise ValueError('Selected epoch lies outside the completed training history')
        if min(row['selected_train_nll'], row['selected_validation_nll']) < 0:
            raise ValueError('Retraining NLL must be nonnegative')
    denominators = report.get('denominators')
    controls = report.get('hyperparameters')
    if not denominators or not controls:
        raise ValueError('Regenerate report with verified denominators and hyperparameters before appending')
    train_n, val_n = denominators['stage_train_images'], denominators['stage_validation_images']
    if any(r['train_images'] != train_n or r['validation_images'] != val_n for r in rows):
        raise ValueError('Retraining image denominators differ across runs')
    clinical_train, clinical_val = denominators['clinical_train_subjects'], denominators['clinical_validation_subjects']
    figure_hashes = report.get('figure_sha256')
    if not isinstance(figure_hashes, dict) or not figure_hashes:
        raise ValueError('Retraining report lacks generated figure hashes; regenerate with plots')
    for name in ('matched_clean_nll.png', 'clinical_clean_nll.png'):
        figure_path = report_path.parent/name
        if not figure_path.is_file() or figure_path.is_symlink():
            raise ValueError(f'Missing retraining figure: {name}')
        if hashlib.sha256(figure_path.read_bytes()).hexdigest() != figure_hashes.get(name):
            raise ValueError(f'Retraining figure hash mismatch: {name}')
    preference = report['validation_only_preference']
    if preference.get('test_used_for_selection') is not False or preference['representative'] not in names:
        raise ValueError('Invalid validation-only retraining preference')
    selected = next(r for r in rows if r['run'] == preference['representative'])
    if selected['mode'] != preference['condition']:
        raise ValueError('Validation representative does not belong to the preferred condition')
    condition = 'stage-only' if preference['condition'] == 'stage_only' else 'multitask'
    checkpoint_checks = report['checkpoint_checks']
    expected_checkpoints = {f'{name}/{checkpoint}.pt' for name in names for checkpoint in ('best', 'last')}
    if len(checkpoint_checks) != 12 or {r['path'] for r in checkpoint_checks} != expected_checkpoints:
        raise ValueError('Retraining checkpoint verification must describe all twelve checkpoint files')
    verified = sum(r['status'] == 'binary_sha256_verified' for r in checkpoint_checks)
    if verified != report['checkpoint_binaries_verified']:
        raise ValueError('Checkpoint verification count is inconsistent')
    checkpoint_total = len(checkpoint_checks)
    report_digest = hashlib.sha256(report_path.read_bytes()).hexdigest()
    blocks = []
    def add(kind, **kwargs): blocks.append(dict(kind=kind, **kwargs))
    def paragraph(text): add('paragraph', text=text)
    def heading(text): add('heading', text=text)
    add('break')
    heading('New retraining with comparable learning curves')
    if synthetic:
        paragraph('SYNTHETIC LOCAL TEST FIXTURE ONLY. These appendix values are not research results.')
    paragraph(f'A separate October 18 experiment completed {report["completed_runs"]} runs and {report["total_epochs"]} epochs, '
              'using stage-only and multitask conditions with three seeds each. At every epoch end, the same fixed checkpoint '
              f'was evaluated on all {train_n:,} training and {val_n:,} validation images with the same clean transform, '
              'evaluation mode and unweighted, unsmoothed negative log-likelihood (NLL). Table 4 and Figure 3 report these new '
              'measurements; Tables 1 to 3 retain the historical October 11 results.')
    display_rows = []
    for row in rows:
        display_rows.append(['Stage only' if row['mode'] == 'stage_only' else 'Multitask', str(row['seed']),
                             str(row['completed_epochs']), str(row['selected_epoch']),
                             f'{row["selected_train_nll"]:.6f}', f'{row["selected_validation_nll"]:.6f}'])
    add('table', caption='Table 4. New retraining epochs and clean NLL at the validation-selected checkpoint.',
        headers=['Condition', 'Seed', 'Completed', 'Selected', 'Train NLL', 'Val NLL'], rows=display_rows,
        widths=[1.02, 1.02, .82, .76, .96, .96])
    add('image', path=str(report_path.parent/'matched_clean_nll.png'), width=6.40,
        caption='Figure 3. New epoch-end clean training and validation NLL for all six runs. Dashed lines identify '
                'validation-selected checkpoints. Both curves use the same measurement procedure; training remains '
                'resubstitution. Raw observations are shown without smoothing or interpolation.')
    add('break')
    heading('Clinical auxiliary curves and interpretation')
    paragraph(f'For the multitask runs, clinical loss was calculated after averaging slice probabilities within each '
              f'subject: {clinical_train} training subjects and {clinical_val} validation subjects. These clean subject-level '
              'curves use the same evaluation procedure across splits. Clinical loss did not select the stage checkpoint.')
    add('image', path=str(report_path.parent/'clinical_clean_nll.png'), width=6.40,
        caption='Figure 4. New clean clinical auxiliary NLL for the three multitask seeds, scored after subject-level '
                'probability aggregation. Dashed lines show the epoch selected by stage-validation NLL. '
                'No held-out or external clinical result is presented.')
    heading('Training controls and validation selection')
    paragraph(f'The verified training configuration used a maximum of {controls["epochs"]} epochs, a minimum of '
              f'{controls["min_epochs"]}, patience of {controls["patience"]} epochs and a required validation-NLL improvement '
              f'greater than {controls["min_delta"]:g}. Inputs were {controls["image_size"]} by {controls["image_size"]} pixels '
              f'with batch size {controls["batch_size"]}. AdamW used learning rate {controls["learning_rate"]:g} and weight decay '
              f'{controls["weight_decay"]:g}; label smoothing was {controls["label_smoothing"]:g}. The multitask auxiliary '
              f'objective had weight {controls["aux_weight"]:g} and was added every {controls["aux_every"]} stage batches.')
    paragraph(f'The validation-only preferred condition was {condition}; its representative used seed {selected["seed"]} '
              f'and epoch {selected["selected_epoch"]}. This preference used mean selected stage-validation NLL across the '
              'three seeds and then the lowest validation NLL within that condition. All six runs remain reported. '
              'Online training NLL measured augmented images with changing parameters; the weighted, smoothed optimization '
              'objective was logged separately and is not substituted for the clean NLL curves.')
    heading('Evidence boundaries and provenance')
    paragraph('Training loss is resubstitution, and validation was used for selection. These new learning curves do not '
              'supply independent test performance, demonstrate clinical superiority or establish readiness for clinical use. '
              'They do not revise the historical benchmark estimates. Image-similarity groups do not establish patient '
              'independence, and the clinical training and validation samples remain small.')
    if verified == checkpoint_total:
        verification_text = f'All {checkpoint_total} retrieved checkpoint file hashes were verified.'
    elif verified:
        verification_text = (f'{verified} of {checkpoint_total} checkpoint file hashes were verified; '
                             'the remaining binaries were absent from this report build.')
    else:
        verification_text = 'The report verified metadata; checkpoint binaries were not supplied for this report build.'
    manifest_checks = report.get('split_checks', {})
    if all(manifest_checks.get(split, {}).get('status') == 'manifest_ids_and_groups_verified'
           for split in ('stage', 'clinical')):
        manifest_text = 'Stage and clinical manifest IDs and group separation were independently checked.'
    else:
        manifest_text = ('Manifest identities are recorded; any manifest not supplied to the report build has '
                         'metadata-only verification.')
    paragraph('The source report checked all six completed histories, per-epoch clean evaluation records, denominators, '
              'stopping decisions and validation-only selection. '+verification_text+
              ' '+manifest_text+' The machine-readable report records source hashes; the appended figures were checked against its figure hashes. '
              f'Report SHA256 begins {report_digest[:16]}; full provenance is retained in report.json and summary.csv.')
    return blocks


def render_docx(blocks, root, output):
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.65)
    sec.left_margin=sec.right_margin=Inches(.75)
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
        style=doc.styles[name]; style.font.name='Calibri'; style.font.color.rgb=RGBColor(0,0,0)
    for element in doc.styles.element.xpath('.//w:pBdr'):
        element.getparent().remove(element)
    doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing=1.06
    doc.styles['Title'].font.size=Pt(21)
    doc.styles['Title'].paragraph_format.space_after=Pt(5)
    doc.styles['Heading 1'].font.size=Pt(13)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(9)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
    doc.styles['Caption'].font.size=Pt(8.5)
    doc.styles['Caption'].font.italic=False
    doc.core_properties.author='Arman Kazi'
    doc.core_properties.title='Dementia MRI Training and Benchmark Results'
    for b in blocks:
        kind=b['kind']
        if kind=='break': doc.add_page_break()
        elif kind in ['title','heading','paragraph','byline']:
            p=doc.add_paragraph(b['text'],style={'title':'Title','heading':'Heading 1'}.get(kind,'Normal'))
            if kind=='byline':
                for run in p.runs: run.font.size=Pt(9)
        elif kind=='table':
            p=doc.add_paragraph(b['caption'],'Caption'); p.paragraph_format.keep_with_next=True
            t=doc.add_table(rows=1,cols=len(b['headers'])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
            for i,h in enumerate(b['headers']): t.rows[0].cells[i].text=h
            for row in b['rows']:
                cells=t.add_row().cells
                for i,v in enumerate(row): cells[i].text=v
            for ri,row in enumerate(t.rows):
                trPr=row._tr.get_or_add_trPr(); no_split=OxmlElement('w:cantSplit');trPr.append(no_split)
                if ri==0: trPr.append(OxmlElement('w:tblHeader'))
                for ci,cell in enumerate(row.cells):
                    cell.width=Inches(b['widths'][ci]*1.18); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcPr=cell._tc.get_or_add_tcPr()
                    borders=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                    tcPr.append(borders)
                    margins=OxmlElement('w:tcMar')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margins.append(e)
                    tcPr.append(margins)
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E3EAF0' if ri==0 else ('F5F6F7' if ri%2==0 else 'FFFFFF'));tcPr.append(shade)
                    for p in cell.paragraphs:
                        p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
                        for run in p.runs: run.font.size=Pt(9);run.bold=(ri==0)
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
        elif kind=='image':
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(root/b['path']),width=Inches(b['width']))
            doc.add_paragraph(b['caption'],'Caption')
        elif kind=='reference':
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(b['url'],RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'1F4E79');rp.append(color);r.append(rp)
            tx=OxmlElement('w:t');tx.text=b['text'];r.append(tx);link.append(r);p._p.append(link)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.save(output)


def markdown(blocks, root, out):
    lines=[]
    for b in blocks:
        k=b['kind']
        if k=='break': continue
        if k=='title': lines+=['# '+b['text'],'']
        elif k=='heading':lines+=['## '+b['text'],'']
        elif k in ['paragraph','byline']:lines += [b['text'],'']
        elif k=='table':
            lines += [b['caption'],'','| '+' | '.join(b['headers'])+' |','| '+' | '.join(['---']*len(b['headers']))+' |']
            lines += ['| '+' | '.join(row)+' |' for row in b['rows']]; lines += ['']
        elif k=='image':
            import os
            path=os.path.relpath(root/b['path'],out.parent)
            lines += [f"![{b['caption']}]({path})",'',b['caption'],'']
        elif k=='reference':lines += [f"- [{b['text']}]({b['url']})"]
    out.write_text('\n'.join(lines).rstrip()+'\n')


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--repo-root',type=Path,default=Path('.'));parser.add_argument('--output-dir',type=Path,default=Path('results/sts/manuscript'));parser.add_argument('--source-commit',required=True)
    parser.add_argument('--retraining-report',type=Path,help='Append two pages from an actual validated completed retraining report.json')
    args=parser.parse_args();root=args.repo_root.resolve();out=args.output_dir.resolve();out.mkdir(parents=True,exist_ok=True)
    blocks,counts,audited=make_content(root,args.source_commit)
    # Regression guard for every displayed supporting numerator, including ensembles.
    lookup={}
    for row in audited:
        lookup[(row.get('cohort'),row.get('model'))]=row
    mapping={'old':'old_baseline','stage_ensemble':'ensemble_stage_only','multitask_ensemble':'ensemble_multitask'}
    cohorts=['internal_same_case','old_hash_group_unexposed_sensitivity','original100']
    for key,values in counts.items():
        for cohort,correct in zip(cohorts,values):
            row=lookup[(cohort,mapping.get(key,key))]
            if int(row['correct'])!=correct: raise ValueError(f'Presentation numerator mismatch: {key} {cohort}')
    if args.retraining_report:
        blocks += make_retraining_appendix(args.retraining_report)
    render_docx(blocks,root,out/'STS_Expanded_Results.docx');markdown(blocks,root,out/'RESULTS.md')
    print(out/'STS_Expanded_Results.docx')
if __name__=='__main__':main()
